"""Attach school quality (kompetenciameres) and school/ovoda provision (KIR) to the spine.

Raw test scores mostly measure the catchment's affluence. The value-added signal is
column 41: result versus what the school's csaladihatter-index predicts. A school
beating its prediction is a school that actually teaches.

Settlement names are unique across Hungary (verified: 0 duplicates in 3178), so the
join is by name.
"""
import json, sys, collections
from pathlib import Path

import atomicjson

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW, OUT = ROOT / "data/raw", ROOT / "data/out"

# OKM 'Adat' columns
O_TYPE, O_GRADE, O_AREA, O_SCORE = 16, 18, 20, 21
O_OM, O_SEQ, O_ZIP, O_TOWN, O_LAT, O_LON = 1, 2, 7, 8, 4, 5
O_NATIONAL, O_CSH, O_VS_CSH = 23, 39, 41
CORE = {"matematika", "szövegértés"}

# KIR columns
K_ZIP, K_TOWN, K_STATUS = 7, 8, 10
K_OVODA, K_ALT1, K_ALT2, K_GIMN = 43, 44, 45, (47, 48, 49)
K_BOLCSODE = 98
K_TORNA, K_USZODA, K_KONYHA, K_KONYV, K_SPORT = 117, 119, 120, 121, 124
K_OV_PLACES, K_ISK_PLACES = 125, 126

OKM_FIELDS = ("school_sites", "school_score_best", "school_score_mean", "school_vs_national",
              "school_beats_prediction", "school_below_prediction")
KIR_FIELDS = ("ovoda_sites", "ovoda_places", "bolcsode_sites", "primary_sites", "primary_places",
              "gimnazium_sites", "has_pool", "has_gym", "has_library", "has_sportsfield")


ROMAN = {v: i + 1 for i, v in enumerate(
    "I II III IV V VI VII VIII IX X XI XII XIII XIV XV XVI XVII XVIII XIX XX XXI XXII XXIII".split())}


def norm_town(t):
    """KIR writes 'Budapest V. kerulet'; the Helysegnevtar writes 'Budapest 05. ker.'."""
    if not t or not t.startswith("Budapest "):
        return t
    r = t.removeprefix("Budapest ").removesuffix(" kerület").rstrip(".")
    return f"Budapest {ROMAN[r]:02d}. ker." if r in ROMAN else t


def bp_district(town, zipcode):
    """OKM (always) and KIR (21 sites) write plain 'Budapest'; the site postcode 1DDx carries the district."""
    z = str(zipcode or "")
    if town != "Budapest" or len(z) != 4 or not z.isdigit() or not 1 <= int(z[1:3]) <= 23:
        return town   # Simplification: 1007 (Margitsziget) has no district digits and stays city-level
    return f"Budapest {int(z[1:3]):02d}. ker."


def sheet(path, name):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return wb, wb[name]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_okm():
    """Per settlement: mean 8th-grade core score, and how sites do vs their CSH prediction."""
    wb, ws = sheet(RAW / "okm_2025.xlsx", "Adat")
    sites = {}
    vs_csh_codes = collections.Counter()
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[O_TYPE] != "általános iskola" or r[O_GRADE] != "8. évfolyam":
            continue
        if r[O_AREA] not in CORE:
            continue
        score = num(r[O_SCORE])
        if score is None:
            continue
        s = sites.setdefault((r[O_OM], r[O_SEQ]), {
            "town": bp_district(r[O_TOWN], r[O_ZIP]), "lat": num(r[O_LAT]), "lon": num(r[O_LON]),
            "scores": [], "national": [], "csh": num(r[O_CSH]), "vs_csh": num(r[O_VS_CSH])})
        s["scores"].append(score)
        s["national"].append(num(r[O_NATIONAL]) or 0)
        vs_csh_codes[r[O_VS_CSH]] += 1
    wb.close()

    by_town = {}
    for s in sites.values():
        s["score"] = sum(s["scores"]) / len(s["scores"])
        s["national_mean"] = sum(s["national"]) / len(s["national"])
        by_town.setdefault(s["town"], []).append(s)
    print(f"  OKM: {len(sites)} primary-school sites in {len(by_town)} settlements")
    print(f"       'vs CSH prediction' code distribution: {dict(vs_csh_codes)}")
    return by_town


def read_kir():
    wb, ws = sheet(RAW / "kir_sites.xlsx", "kir_felhelyek_2026_08_30")
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[K_STATUS] != "működő":
            continue
        t = out.setdefault(bp_district(norm_town(r[K_TOWN]), r[K_ZIP]), collections.Counter())
        flag = lambda i: 1 if r[i] == "X" else 0
        yn = lambda i: 1 if r[i] == "I" else 0
        t["ovoda"] += flag(K_OVODA)
        t["primary"] += 1 if (flag(K_ALT1) or flag(K_ALT2)) else 0
        t["gimnazium"] += 1 if any(flag(i) for i in K_GIMN) else 0
        t["bolcsode"] += flag(K_BOLCSODE)
        t["ovoda_places"] += int(num(r[K_OV_PLACES]) or 0)
        t["primary_places"] += int(num(r[K_ISK_PLACES]) or 0)
        for key, col in (("torna", K_TORNA), ("uszoda", K_USZODA), ("konyha", K_KONYHA),
                         ("konyvtar", K_KONYV), ("sportpalya", K_SPORT)):
            t[key] += yn(col)
    wb.close()
    print(f"  KIR: {len(out)} settlements with at least one operating site")
    return out


def main():
    okm, kir = read_okm(), read_kir()
    spine = json.loads((OUT / "settlements.json").read_text(encoding="utf-8"))

    hit_okm = hit_kir = 0
    for s in spine:
        n = s["name"]
        sites = okm.get(n)
        # a settlement that lost its match must lose the old values, not keep a stale copy
        for key in (() if sites else OKM_FIELDS) + (() if kir.get(n) else KIR_FIELDS):
            s.pop(key, None)
        if sites:
            hit_okm += 1
            best = max(sites, key=lambda x: x["score"])
            s["school_sites"] = len(sites)
            s["school_score_best"] = round(best["score"])
            s["school_score_mean"] = round(sum(x["score"] for x in sites) / len(sites))
            s["school_vs_national"] = round(s["school_score_mean"] - best["national_mean"])
            # code 1 = above CSH prediction, 0 = as predicted, -1 = below, 6 = no index
            graded = [x["vs_csh"] for x in sites if x["vs_csh"] in (-1, 0, 1)]  # 2-4,6,7 = no comparison
            s["school_beats_prediction"] = (
                round(100 * sum(1 for v in graded if v == 1) / len(graded)) if graded else None)
            s["school_below_prediction"] = (
                round(100 * sum(1 for v in graded if v == -1) / len(graded)) if graded else None)
        k = kir.get(n)
        if k:
            hit_kir += 1
            s["ovoda_sites"] = k["ovoda"]
            s["ovoda_places"] = k["ovoda_places"]
            s["bolcsode_sites"] = k["bolcsode"]
            s["primary_sites"] = k["primary"]
            s["primary_places"] = k["primary_places"]
            s["gimnazium_sites"] = k["gimnazium"]
            s["has_pool"] = bool(k["uszoda"])
            s["has_gym"] = bool(k["torna"])
            s["has_library"] = bool(k["konyvtar"])
            s["has_sportsfield"] = bool(k["sportpalya"])

    atomicjson.write_json(OUT / "settlements.json", spine, ensure_ascii=False, indent=1)
    print(f"\nsettlements with school results : {hit_okm}")
    print(f"settlements with KIR provision  : {hit_kir}")
    print(f"  with an ovoda                 : {sum(1 for s in spine if s.get('ovoda_sites'))}")
    print(f"  with a primary school         : {sum(1 for s in spine if s.get('primary_sites'))}")
    return spine


def demo():
    spine = main()
    by = {s["name"]: s for s in spine}
    assert by["Gödöllő"].get("school_sites"), "Godollo must have measured schools"
    assert 1400 < by["Gödöllő"]["school_score_mean"] < 1900, by["Gödöllő"]["school_score_mean"]
    assert by["Budapest 13. ker."]["ovoda_sites"] > 10, "district ovodas must join"
    assert bp_district("Budapest", "1078") == "Budapest 07. ker." and bp_district("Vác", "2600") == "Vác"
    assert bp_district("Budapest", 1039) == "Budapest 03. ker."   # KIR stores the postcode as int
    assert by["Budapest"].get("school_sites") is None, "city-wide row must not double-count districts"
    assert all(by[f"Budapest {d:02d}. ker."].get("school_sites") for d in range(1, 24)), \
        "every district must get OKM results via its postcode"
    assert by["Gödöllő"]["ovoda_sites"] > 0 and by["Gödöllő"]["primary_sites"] > 0
    scored = [s for s in spine if s.get("school_score_mean")]
    assert len(scored) > 900, f"only {len(scored)} settlements with scores"
    for s in scored:
        assert 1000 < s["school_score_mean"] < 2200, (s["name"], s["school_score_mean"])
    ov = [s for s in spine if s.get("ovoda_sites")]
    assert len(ov) > 1800, f"only {len(ov)} settlements with an ovoda"
    print("self-check OK")


if __name__ == "__main__":
    demo() if "--check" in sys.argv else main()
